"""Dedicated PPTX reader with conversion-local placeholder geometry caching.

Rendering rules adapted from Microsoft MarkItDown 0.1.7 (MIT, see NOTICE).
Model enrichment remains in Markitai's workflow rather than this file reader.
"""

from __future__ import annotations

import base64
import html
import re
from pathlib import Path
from typing import Any, cast

from markitai.converter.base import ConvertResult
from markitai.converter.structured_text import _normalize


class ShapePositions:
    """Resolve both coordinates together, caching only this immutable document.

    python-pptx exposes inherited placeholder geometry as properties. Sorting
    asks for those repeatedly and re-scans layout/master XML on each access.
    This cache is discarded after conversion; it never changes library classes
    or shares geometry between files or concurrent conversions.
    """

    def __init__(self) -> None:
        self._positions: dict[Any, tuple[Any, Any]] = {}
        self._layout_bases: dict[tuple[Any, int], Any] = {}

    def _resolve(self, shape: Any) -> tuple[Any, Any]:
        from pptx.shapes.placeholder import (
            PlaceholderPicture,
            _BaseSlidePlaceholder,
            _InheritsDimensions,
        )

        element = shape._element
        if element in self._positions:
            return self._positions[element]
        if isinstance(shape, _InheritsDimensions):
            top, left = element.y, element.x
            if top is None or left is None:
                if isinstance(shape, (_BaseSlidePlaceholder, PlaceholderPicture)):
                    key = (cast(Any, shape.part).slide_layout.part, element.ph_idx)
                    if key not in self._layout_bases:
                        self._layout_bases[key] = shape._base_placeholder
                    base = self._layout_bases[key]
                else:
                    base = shape._base_placeholder
                if base is not None:
                    base_top, base_left = self._resolve(base)
                    top = base_top if top is None else top
                    left = base_left if left is None else left
        else:
            top, left = shape.top, shape.left
        self._positions[element] = top, left
        return top, left

    def key(self, shape: Any) -> tuple[Any, Any]:
        top, left = self._resolve(shape)
        # Preserve the reference's ordering of missing AND zero coordinates.
        return top or float("-inf"), left or float("-inf")


def _svg_part(shape: Any) -> Any:
    try:
        namespace = {
            "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
        }
        embed = (
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        )
        for blip in shape._element.findall(".//asvg:svgBlip", namespace):
            rid = blip.get(embed)
            if rid:
                return shape.part.related_part(rid)
    except Exception:
        pass
    return None


def _image_data(shape: Any) -> tuple[bytes | None, str | None]:
    try:
        image = shape.image
        return image.blob, image.content_type
    except Exception:
        part = _svg_part(shape)
        if part is not None:
            try:
                return part.blob, "image/svg+xml"
            except Exception:
                pass
    return None, None


def _is_picture(shape: Any, kind: Any) -> bool:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if kind == MSO_SHAPE_TYPE.PICTURE:
        return True
    if kind == MSO_SHAPE_TYPE.PLACEHOLDER:
        try:
            return shape.image is not None
        except Exception:
            return _svg_part(shape) is not None
    return False


def _picture(shape: Any) -> str:
    blob, content_type = _image_data(shape)
    alt = ""
    try:
        alt = shape._element._nvXxPr.cNvPr.attrib.get("descr", "")
    except Exception:
        pass
    alt = re.sub(r"\s+", " ", re.sub(r"[\r\n\[\]]", " ", alt)).strip()
    if blob is not None:
        encoded = base64.b64encode(blob).decode("utf-8")
        return f"\n![{alt}](data:{content_type or 'image/png'};base64,{encoded})\n"
    name = re.sub(r"\W", "", shape.name) + ".jpg"
    return f"\n![{alt}]({name})\n"


def _table(table: Any) -> str:
    from markitai.converter.xlsx import dataframe_table_markdown

    rows = []
    for index, row in enumerate(table.rows):
        tag = "th" if index == 0 else "td"
        cells = "".join(
            f"<{tag}>{html.escape(cell.text)}</{tag}>" for cell in row.cells
        )
        rows.append(f"<tr>{cells}</tr>")
    if not rows:
        return "\n"
    source = (
        f"<table><thead>{rows[0]}</thead><tbody>{''.join(rows[1:])}</tbody></table>"
    )
    return dataframe_table_markdown(source) + "\n"


def _chart(chart: Any) -> str:
    try:
        heading = "\n\n### Chart"
        if chart.has_title:
            heading += f": {chart.chart_title.text_frame.text}"
        categories = [category.label for category in chart.plots[0].categories]
        series = list(chart.series)
        values = [list(item.values) for item in series]
        rows = [["Category", *(item.name for item in series)]]
        for index, name in enumerate(categories):
            rows.append([name, *(v[index] if index < len(v) else None for v in values)])
        lines = ["| " + " | ".join(map(str, row)) + " |" for row in rows]
        separator = "|" + "|".join(["---"] * len(rows[0])) + "|"
        return heading + "\n\n" + "\n".join([lines[0], separator, *lines[1:]])
    except ValueError as exc:
        if "unsupported plot type" not in str(exc):
            raise
        return "\n\n[unsupported chart]\n\n"
    except Exception:
        return "\n\n[unsupported chart]\n\n"


def _render_shape(
    shape: Any, title: Any, positions: ShapePositions, chunks: list[str]
) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    kind = shape.shape_type
    if _is_picture(shape, kind):
        chunks.append(_picture(shape))
    if kind == MSO_SHAPE_TYPE.TABLE:
        chunks.append(_table(shape.table))
    if shape.has_chart:
        chunks.append(_chart(shape.chart))
    elif shape.has_text_frame:
        chunks.append(
            ("# " + shape.text.lstrip() if shape == title else shape.text) + "\n"
        )
    if kind == MSO_SHAPE_TYPE.GROUP:
        for child in sorted(shape.shapes, key=positions.key):
            _render_shape(child, title, positions, chunks)


def convert_pptx(path: Path) -> ConvertResult:
    from pptx import Presentation

    deck = Presentation(str(path))
    positions = ShapePositions()
    slides = []
    for number, slide in enumerate(deck.slides, 1):
        title = slide.shapes.title
        chunks = [f"<!-- Slide number: {number} -->\n"]

        for shape in sorted(slide.shapes, key=positions.key):
            _render_shape(shape, title, positions, chunks)
        text = "".join(chunks).strip()
        if slide.has_notes_slide:
            text += "\n\n### Notes:\n"
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                text += frame.text
        slides.append(text.strip())
    return ConvertResult(
        markdown=_normalize("\n\n".join(slides).strip()),
        metadata={"source": str(path), "format": "PPTX", "converter": "python-pptx"},
    )
