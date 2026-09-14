"""Plain tables may avoid pandas startup; complex cells retain its semantics."""

import datetime
import json
import os
import subprocess
import sys
from typing import Any, cast

import pandas as pd
import pytest
from pandas._libs.parsers import STR_NA_VALUES


@pytest.mark.parametrize(
    "rows",
    [
        [],
        [["name", "value"]],
        [["name", "value"], ["中文<&>*_|", 42], ["other", -5]],
        [
            [0, "date text", "enabled"],
            [1, "15/10/2017", True],
            [2, "16/08/2016", False],
        ],
        [["text"], ["a" * 1000], [" a b "], ["&amp; <script>"]],
    ],
)
def test_plain_table_matches_reader_and_formatter(rows, tmp_path):
    from openpyxl import Workbook

    from markitai.converter.xlsx import dataframe_markdown, plain_table_markdown

    path = tmp_path / "cells.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    book.save(path)
    expected = dataframe_markdown(pd.read_excel(path, engine="openpyxl"))
    assert plain_table_markdown(rows) == expected


@pytest.mark.parametrize(
    "rows",
    [
        [["x"], [None]],
        [["x"], ["NA"]],
        [["x"], ["True"]],
        [["x"], ["001"]],
        [["x"], [1.25]],
        [["x"], [datetime.date(2024, 1, 1)]],
        [["x"], ["a  b"]],
        [["x"], ["a\nb"]],
        [["x"], ["a\tb"]],
        [["x"], [2**64]],
        [["x"], [1], ["mixed"]],
        [["dup", "dup"], [1, 2]],
        [[None], [1]],
        [["x", "y"], [1]],
    ],
)
def test_ambiguous_formatting_uses_pandas(rows):
    from markitai.converter.xlsx import plain_table_markdown

    assert plain_table_markdown(rows) is None


def test_plain_workbook_cold_conversion_does_not_import_pandas(tmp_path):
    from openpyxl import Workbook

    from markitai.converter.office import XlsxConverter

    path = tmp_path / "ordinary.xlsx"
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row in [
        [0, "name", "date"],
        [1, "中文", "15/10/2017"],
        [2, "hello", "16/08/2016"],
    ]:
        sheet.append(row)
    book.save(path)
    expected = XlsxConverter().convert(path).markdown
    code = """import json,sys
from pathlib import Path
from markitai.converter.office import XlsxConverter
r = XlsxConverter().convert(Path(sys.argv[1]))
assert "pandas" not in sys.modules
print(json.dumps({"markdown": r.markdown, "metadata": r.metadata}))
"""
    result = subprocess.run(
        [sys.executable, "-c", code, str(path)],
        capture_output=True,
        text=True,
        env={**os.environ, "LITELLM_LOCAL_MODEL_COST_MAP": "True"},
    )
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["markdown"] == expected


@pytest.mark.parametrize("value", sorted(STR_NA_VALUES))
def test_every_pandas_na_token_uses_reference_parser(value):
    from markitai.converter.xlsx import plain_table_markdown

    assert plain_table_markdown([["value"], [value]]) is None


def test_generated_plain_rows_match_pandas_inference():
    import random

    from pandas.io.parsers import TextParser

    from markitai.converter.xlsx import dataframe_markdown, plain_table_markdown

    rng = random.Random(20260926)
    alphabet = "abc中é<&>_;'\"*`|\\ \u00a0\u2003"
    checked = 0
    for _ in range(1000):
        rows = [[0, "text", "boolean"]] + [
            [
                rng.randrange(-(2**40), 2**40),
                "".join(rng.choices(alphabet, k=20)),
                bool(rng.randrange(2)),
            ]
            for _ in range(8)
        ]
        actual = plain_table_markdown(rows)
        if actual is not None:
            checked += 1
            assert actual == dataframe_markdown(
                TextParser(rows, header=0, skip_blank_lines=False).read()
            ), rows
    assert checked > 500


def test_complete_reader_matrix_keeps_typed_frames(tmp_path):
    import importlib.util
    from pathlib import Path

    from openpyxl import load_workbook

    from markitai.converter.xlsx import _pandas_sheet, _sheet_rows

    script = (
        Path(__file__).resolve().parents[4] / "scripts/benchmarks/excel_fixtures.py"
    )
    spec = importlib.util.spec_from_file_location("excel_fixtures", script)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    for path in module.fixtures(tmp_path):
        expected = pd.read_excel(path, engine="openpyxl", sheet_name=None)
        with_book = load_workbook(
            path, read_only=True, data_only=True, keep_links=False
        )
        try:
            for sheet in with_book.worksheets:
                actual = _pandas_sheet(_sheet_rows(sheet), sheet.title)
                pd.testing.assert_frame_equal(actual, expected[sheet.title])
        finally:
            with_book.close()


def test_fallback_reads_each_worksheet_only_once(tmp_path, monkeypatch):
    from collections import Counter

    from openpyxl import Workbook
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    from markitai.converter.xlsx import convert_xlsx

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["float"])
    sheet.append([1.25])
    book.create_sheet("Empty")
    path = tmp_path / "fallback.xlsx"
    book.save(path)
    calls = Counter()
    original = cast(Any, ReadOnlyWorksheet)._cells_by_row

    def counted(self, *args, **kwargs):
        calls[self.title] += 1
        return original(self, *args, **kwargs)

    monkeypatch.setattr(ReadOnlyWorksheet, "_cells_by_row", counted)
    convert_xlsx(path)
    assert calls == {"Sheet": 1, "Empty": 1}
