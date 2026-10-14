"""Spreadsheets as Markdown: empty cells empty, banner rows as paragraphs.

The fixtures are invented: a fair's day plan laid out the way hand-made
schedules are (a merged title and contact lines above the table, a merged
section label inside it, a column whose header cell is empty).
"""

from __future__ import annotations

import datetime
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from markitai.converter.office import XlsConverter, XlsxConverter
from markitai.converter.opendocument import OdsConverter
from markitai.converter.sheet_tables import banner_text, sheet_markdown

FIXTURES = Path(__file__).parent.parent / "fixtures"

DAY_PLAN = [
    ["Spring Fair 2026 · Day Plan", None, None, None],
    ["Venue: Riverside Hall, east wing", None, None, None],
    ["Item", "Time", "Task", "Note"],
    ["Setup crew", "07:30", "Unload tables", None],
    ["Stage", None, "Sound check", None],
    ["Afternoon", None, None, None],
    [None, "Time", "Task", "Note"],
    ["Raffle", "14:00", "Draw tickets", "NA"],
    ["Close", 17.5, None, 3],
]
BANNER_ROWS = (1, 2, 6)  # 1-based rows merged across A:D

DAY_PLAN_MARKDOWN = """## Day plan

Spring Fair 2026 · Day Plan

Venue: Riverside Hall, east wing

| Item | Time | Task | Note |
| --- | --- | --- | --- |
| Setup crew | 07:30 | Unload tables |  |
| Stage |  | Sound check |  |

Afternoon

|  | Time | Task | Note |
| --- | --- | --- | --- |
| Raffle | 14:00 | Draw tickets | NA |
| Close | 17.5 |  | 3 |"""


def _workbook(
    path: Path, rows: list[list[object]], merged=(), title="Day plan"
) -> Path:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = title
    for row in rows:
        sheet.append(row)
    for row in merged:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    book.save(path)
    return path


class TestXlsx:
    def test_merged_rows_are_paragraphs_and_empty_cells_stay_empty(self, tmp_path):
        path = _workbook(tmp_path / "plan.xlsx", DAY_PLAN, BANNER_ROWS)
        markdown = XlsxConverter().convert(path).markdown
        assert markdown == DAY_PLAN_MARKDOWN
        assert "NaN" not in markdown
        assert "Unnamed" not in markdown

    def test_an_unmerged_single_value_row_stays_in_the_table(self, tmp_path):
        # Only a merge marks a banner: a row with one value is data
        path = _workbook(tmp_path / "plain.xlsx", DAY_PLAN)
        markdown = XlsxConverter().convert(path).markdown
        assert "| Spring Fair 2026 · Day Plan |  |  |  |\n| --- |" in markdown
        assert "| Afternoon |  |  |  |" in markdown

    def test_cells_show_their_own_text(self, tmp_path):
        rows = [
            ["name", "flag", "ratio", "when", "sum", None],
            ["a", True, 0.1, datetime.datetime(2024, 1, 2), "NA", None],
            ["b", False, 1e-9, datetime.datetime(2024, 1, 2, 13, 5), 12, None],
            ["c*_|x", None, 1234567.125, datetime.date(2024, 3, 4), None, "x"],
        ]
        path = _workbook(tmp_path / "typed.xlsx", rows, title="Typed")
        assert XlsxConverter().convert(path).markdown == (
            "## Typed\n\n"
            "| name | flag | ratio | when | sum |  |\n"
            "| --- | --- | --- | --- | --- | --- |\n"
            "| a | TRUE | 0.1 | 2024-01-02 | NA |  |\n"
            "| b | FALSE | 0.000000001 | 2024-01-02 13:05:00 | 12 |  |\n"
            "| c*_\\|x |  | 1234567.125 | 2024-03-04 |  | x |"
        )

    def test_gaps_inside_a_table_keep_their_rows(self, tmp_path):
        rows = [[None], ["A", "B", "C", "D"], [1, None, None, 4], [None], [5], [None]]
        path = _workbook(tmp_path / "gaps.xlsx", rows, title="Gaps")
        assert XlsxConverter().convert(path).markdown == (
            "## Gaps\n\n"
            "| A | B | C | D |\n"
            "| --- | --- | --- | --- |\n"
            "| 1 |  |  | 4 |\n"
            "|  |  |  |  |\n"
            "| 5 |  |  |  |"
        )

    def test_error_cells_show_their_code(self, tmp_path):
        from openpyxl import Workbook

        path = tmp_path / "errors.xlsx"
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.append(["a", "b"])
        sheet.append([1, 2])
        book.save(path)
        # openpyxl cannot write a cached error value; patch the sheet part
        with zipfile.ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        sheet_xml = parts["xl/worksheets/sheet1.xml"].decode()
        sheet_xml = sheet_xml.replace(
            '<c r="B2" t="n"><v>2</v></c>', '<c r="B2" t="e"><v>#DIV/0!</v></c>'
        )
        assert "#DIV/0!" in sheet_xml
        parts["xl/worksheets/sheet1.xml"] = sheet_xml.encode()
        with zipfile.ZipFile(path, "w") as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
        assert "| 1 | #DIV/0! |" in XlsxConverter().convert(path).markdown

    def test_conversion_does_not_load_pandas(self, tmp_path):
        path = _workbook(tmp_path / "plan.xlsx", DAY_PLAN, BANNER_ROWS)
        code = (
            "import sys\nfrom pathlib import Path\n"
            "from markitai.converter.office import XlsxConverter\n"
            "md = XlsxConverter().convert(Path(sys.argv[1])).markdown\n"
            "assert 'Afternoon' in md, md\n"
            "assert 'pandas' not in sys.modules\n"
        )
        result = subprocess.run(
            [sys.executable, "-c", code, str(path)],
            capture_output=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
        )
        assert result.returncode == 0, result.stderr


class TestXls:
    def test_fixture_reads_merges_like_xlsx(self):
        # tests/fixtures/day-plan.xls: written with xlwt (merged rows 1 and 5)
        result = XlsConverter().convert(FIXTURES / "day-plan.xls")
        assert result.metadata["converter"] == "xlrd"
        assert result.markdown == (
            "## Day plan\n\n"
            "Spring Fair 2026 · Day Plan\n\n"
            "| Item | Time | Task | Note |\n"
            "| --- | --- | --- | --- |\n"
            "| Setup crew | 07:30 | Unload tables |  |\n"
            "| Stage |  | Sound check |  |\n\n"
            "Afternoon\n\n"
            "|  | Time | Task | Note |\n"
            "| --- | --- | --- | --- |\n"
            "| Raffle | 14:00 | Draw tickets | NA |\n"
            "| Close | 17.5 |  | 3 |"
        )

    def test_not_an_xls_falls_back(self, tmp_path):
        path = tmp_path / "notes.xls"
        path.write_text("plain notes, not a workbook\n", encoding="utf-8")
        assert "plain notes" in XlsConverter().convert(path).markdown


class TestOds:
    def test_spanned_cells_are_banner_rows(self, tmp_path):
        ns = (
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
            'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"'
        )

        def cell(text: str = "", span: int = 1) -> str:
            spanned = f' table:number-columns-spanned="{span}"' if span > 1 else ""
            body = f"<text:p>{text}</text:p>" if text else ""
            return f"<table:table-cell{spanned}>{body}</table:table-cell>"

        def row(*cells: str) -> str:
            return "<table:table-row>" + "".join(cells) + "</table:table-row>"

        covered = "<table:covered-table-cell/>" * 2
        body = (
            '<office:spreadsheet><table:table table:name="Plan">'
            + row(cell("Spring Fair 2026", 3), covered)
            + row(cell("Item"), cell("Time"), cell("Task"))
            + row(cell("Stage"), cell(), cell("Sound check"))
            + row(cell("Afternoon", 3), covered)
            + row(cell(), cell("Time"), cell("Task"))
            + row(cell("Raffle"), cell("14:00"), cell("Draw tickets"))
            + "</table:table></office:spreadsheet>"
        )
        path = tmp_path / "plan.ods"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr(
                "content.xml",
                f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ns}>'
                f"<office:body>{body}</office:body></office:document-content>",
            )
        assert OdsConverter().convert(path).markdown == (
            "## Plan\n\n"
            "Spring Fair 2026\n\n"
            "| Item | Time | Task |\n"
            "| --- | --- | --- |\n"
            "| Stage |  | Sound check |\n\n"
            "Afternoon\n\n"
            "|  | Time | Task |\n"
            "| --- | --- | --- |\n"
            "| Raffle | 14:00 | Draw tickets |\n"
        )


class TestSheetMarkdown:
    def test_a_banner_needs_a_merge_of_two_columns_or_more(self):
        rows = [["Title", ""], ["a", "b"], ["1", "2"]]
        assert sheet_markdown(rows, {(0, 0): 1}).startswith("| Title |  |")
        assert sheet_markdown(rows, {(0, 0): 2}).startswith("Title\n\n| a | b |")

    def test_columns_empty_in_a_block_are_left_out(self):
        rows = [["Head", "", "x"], ["a", "", "b"], ["Part two", "", ""], ["c", "", ""]]
        assert sheet_markdown(rows, {(2, 0): 3}) == (
            "| Head | x |\n| --- | --- |\n| a | b |\n\nPart two\n\n| c |\n| --- |"
        )

    @pytest.mark.parametrize(
        ("text", "expected"),
        [
            ("# not a heading", "\\# not a heading"),
            ("- not a list", "\\- not a list"),
            ("1. not a list", "1\\. not a list"),
            ("2) nor this", "2\\) nor this"),
            ("Draft *final*", "Draft \\*final\\*"),
            ("Line one\nline two", "Line one line two"),
        ],
    )
    def test_banner_text_is_literal(self, text, expected):
        assert banner_text(text) == expected
