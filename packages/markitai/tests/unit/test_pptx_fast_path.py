"""Real PPTX documents constrain the native reader and geometry optimization."""

from typing import Any, cast
from unittest.mock import patch

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from markitai.converter.office import OfficeConverter, PptxConverter


@pytest.fixture
def rich_deck(tmp_path, fixtures_dir):
    deck = cast(Any, Presentation())
    for layout in [0, 1, 2, 3, 5, 6]:
        slide = deck.slides.add_slide(deck.slide_layouts[layout])
        if slide.shapes.title is not None:
            slide.shapes.title.text = f"Title {layout}"
        for shape in slide.placeholders:
            if shape != slide.shapes.title and shape.has_text_frame:
                shape.text = f"Body {shape.placeholder_format.idx}"
        slide.notes_slide.notes_text_frame.text = f"Speaker notes {layout}"
    slide = deck.slides[-1]
    group = slide.shapes.add_group_shape()
    for top, left, text in [(2, 1, "Second"), (1, 0, "First"), (0, 0, "Zero")]:
        group.shapes.add_textbox(
            Inches(left), Inches(top), Inches(2), Inches(1)
        ).text = text
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(4), Inches(4), Inches(1)
    ).table
    for row, values in zip(
        table.rows, [["*Header*", "中文"], ["Line\nbreak", "<x> & a|b"]]
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    slide.shapes.add_picture(
        str(fixtures_dir / "sample.jpg"), Inches(4), Inches(1), width=Inches(1)
    )
    chart = CategoryChartData()
    chart.categories = ["A", "B"]
    chart.add_series("Values", [1, 2])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(5),
        Inches(4),
        Inches(2),
        Inches(2),
        chart,
    )
    path = tmp_path / "rich.pptx"
    deck.save(path)
    return path


def test_rich_pptx_keeps_reference_output(rich_deck):
    actual = PptxConverter().convert(rich_deck).markdown
    assert actual == OfficeConverter().convert(rich_deck).markdown
    for text in [
        "Speaker notes",
        "Chart",
        "data:image/jpeg;base64,",
        "\\*Header\\*",
        "First",
    ]:
        assert text in actual


def test_placeholder_geometry_is_resolved_once_per_base(rich_deck):
    from pptx.shapes.placeholder import SlidePlaceholder

    from markitai.converter.pptx import ShapePositions

    deck = cast(Any, Presentation(rich_deck))
    positions = ShapePositions()
    shape = deck.slides[1].shapes.title
    expected = (shape.top or float("-inf"), shape.left or float("-inf"))
    getter = SlidePlaceholder._base_placeholder.fget
    assert getter is not None
    calls = []

    def tracked(self):
        calls.append(self)
        return getter(self)

    with patch.object(SlidePlaceholder, "_base_placeholder", property(tracked)):
        assert positions.key(shape) == expected
        assert positions.key(shape) == expected
    assert len(calls) == 1


def test_shared_layout_lookup_is_reused_across_slides():
    from pptx.shapes.placeholder import SlidePlaceholder

    from markitai.converter.pptx import ShapePositions

    deck = cast(Any, Presentation())
    shapes = [
        deck.slides.add_slide(deck.slide_layouts[1]).shapes.title for _ in range(20)
    ]
    expected = [
        (shape.top or float("-inf"), shape.left or float("-inf")) for shape in shapes
    ]
    getter = SlidePlaceholder._base_placeholder.fget
    assert getter is not None
    calls = []

    def tracked(self):
        calls.append(self)
        return getter(self)

    with patch.object(SlidePlaceholder, "_base_placeholder", property(tracked)):
        positions = ShapePositions()
        assert [positions.key(shape) for shape in shapes] == expected
    assert len(calls) == 1


def test_mislabeled_input_uses_generic_detection(fixtures_dir, tmp_path):
    path = tmp_path / "actually-word.pptx"
    path.write_bytes((fixtures_dir / "sample.docx").read_bytes())
    assert (
        PptxConverter().convert(path).markdown
        == OfficeConverter().convert(path).markdown
    )


def test_plain_text_mislabeled_as_pptx_uses_generic_detection(tmp_path):
    path = tmp_path / "text.pptx"
    path.write_text("A complete plain text document.")
    assert (
        PptxConverter().convert(path).markdown
        == OfficeConverter().convert(path).markdown
    )


def test_svg_picture_without_raster_fallback(rich_deck, tmp_path):
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part
    from pptx.opc.packuri import PackURI
    from pptx.oxml.ns import qn

    deck = cast(Any, Presentation(rich_deck))
    slide = deck.slides[-1]
    picture = next(s for s in slide.shapes if s.shape_type == 13)
    blob = b'<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20"><rect width="20" height="20"/></svg>'
    part = Part(
        PackURI("/ppt/media/vector.svg"), "image/svg+xml", slide.part.package, blob
    )
    rid = slide.part.relate_to(part, RT.IMAGE)
    blip = picture._element.find(".//" + qn("a:blip"))
    assert blip is not None
    blip.attrib.pop(qn("r:embed"))
    ext_list = etree.SubElement(blip, qn("a:extLst"))
    ext = etree.SubElement(
        ext_list, qn("a:ext"), uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
    )
    svg = etree.SubElement(
        ext, "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip"
    )
    svg.set(qn("r:embed"), rid)
    path = tmp_path / "vector.pptx"
    deck.save(path)
    actual = PptxConverter().convert(path).markdown
    assert "data:image/svg+xml;base64," in actual
    assert actual == OfficeConverter().convert(path).markdown
