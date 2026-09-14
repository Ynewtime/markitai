"""Synthetic workbooks shared by correctness and performance audits."""

# Excel stores naive dates and times; retain them as fixture inputs.
# ruff: noqa: DTZ001

from __future__ import annotations

import datetime as dt
import random
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.datetime import CALENDAR_MAC_1904


def write_book(path: Path, rows: list, **options) -> None:
    book = Workbook()
    if options.get("epoch1904"):
        book.epoch = CALENDAR_MAC_1904
    book.iso_dates = options.get("iso_dates", False)
    sheet = book.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    if options.get("merge"):
        sheet.merge_cells("A1:C1")
    if options.get("styled_tail"):
        sheet["G40"].font = Font(bold=True)
    book.create_sheet("Empty")
    hidden = book.create_sheet("隐藏")
    hidden.sheet_state = "hidden"
    hidden.append(["name", "value"])
    hidden.append(["中文", 4])
    book.save(path)


def replace_member(path: Path, transform) -> None:
    with zipfile.ZipFile(path) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with zipfile.ZipFile(path, "w") as archive:
        for item, data in members:
            if item.filename == "xl/worksheets/sheet1.xml":
                data = transform(data.decode()).encode()
            archive.writestr(item, data)


def fixtures(directory: Path) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    cases = {
        "plain": [["name", "value"], ["中文", 42], [None, False]],
        "headers": [["dup", "dup", None], [1, 2, 3], [None, None, None]],
        "leading_empty": [[], [], [None, None, "name", "value"], [None, None, "x", 42]],
        "na_strings": [
            ["text", "other"],
            ["NA", "NULL"],
            ["NaN", "001"],
            ["n/a", "inf"],
        ],
        "escape_strings": [
            ["text"],
            ["_x0041_"],
            ["_x005F_x0041_"],
            ["a\r\nb\t c"],
            ["&amp; <b>*x*</b>"],
        ],
        "errors": [["error", "value"], ["#DIV/0!", 1], ["#N/A", 2], ["#VALUE!", 3]],
        "booleans": [
            [True, False, "mixed"],
            [False, True, 1],
            [1, 0, False],
            [None, "x", "TRUE"],
        ],
        "floats": [
            ["n"],
            [1.000000000000001],
            [1e-20],
            [-0.0],
            [1e30],
            [1.2345678901234567],
        ],
        "dates": [
            ["date", "mixed"],
            [dt.datetime(2024, 1, 2, 3, 4, 5, 123456), "text"],
            [dt.datetime(1900, 2, 28), dt.datetime(2020, 1, 1)],
            [None, "tail"],
        ],
        "time": [
            ["time", "mixed"],
            [dt.time(23, 59, 59, 123456), "text"],
            [dt.time(0, 0, 0, 999), dt.time(12)],
            [None, "tail"],
        ],
        "duration": [
            ["duration", "mixed"],
            [dt.timedelta(days=3, microseconds=123456), "text"],
            [dt.timedelta(seconds=-5), dt.timedelta(hours=7)],
            [None, "tail"],
        ],
        "formula": [["formula", "value"], ["=1+2", 3], ["=A2*2", 6]],
        "empty": [],
        "merge": [["merged", None, None], [1, 2, 3]],
        "styled_tail": [["x", "y"], [1, 2]],
    }
    for name, rows in cases.items():
        write_book(directory / f"{name}.xlsx", rows, **{name: True})
    for option in ("epoch1904", "iso_dates"):
        write_book(directory / f"{option}.xlsx", cases["dates"], **{option: True})
    rng = random.Random(20260914)
    alphabet = "abc中é<&>;'\"_*`|\\\r\n\t \u00a0"
    generated = [[f"column {i}" for i in range(8)]] + [
        ["".join(rng.choices(alphabet, k=30)) for _ in range(8)] for _ in range(100)
    ]
    write_book(directory / "generated_strings.xlsx", generated)
    # Keep XML integer tokens exact: Excel writers may already round them.
    numbers = [
        str(2**53 - 1),
        str(2**53 + 1),
        str(2**63 - 1),
        str(2**64 - 1),
        str(-(2**53 + 1)),
    ]
    path = directory / "large_integer_tokens.xlsx"
    write_book(path, [["value"]] + [[i] for i in range(len(numbers))])
    iterator = iter(numbers)
    replace_member(
        path,
        lambda xml: re.sub(r"<v>\d+</v>", lambda _: f"<v>{next(iterator)}</v>", xml),
    )
    path = directory / "cached_formula.xlsx"
    write_book(path, cases["formula"])
    replace_member(
        path,
        lambda xml: xml.replace("<f>1+2</f><v></v>", "<f>1+2</f><v>3</v>").replace(
            "<f>A2*2</f><v></v>", "<f>A2*2</f><v>6</v>"
        ),
    )
    path = directory / "wrong_dimension.xlsx"
    write_book(path, cases["plain"])
    replace_member(
        path,
        lambda xml: re.sub(r'<dimension ref="[^"]+"', '<dimension ref="A1:A1"', xml),
    )
    return sorted(directory.glob("*.xlsx"))
