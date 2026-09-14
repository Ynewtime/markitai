"""Paired Office size scaling, with cold processes and exact Markdown hashes.

The original XLSX invocation remains the default; PPTX and DOCX measure slides
and paragraphs respectively. Additional sizes remain diagnostic workloads.
"""

from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
import time
import zipfile
from pathlib import Path
from typing import Any, cast

REPO = Path(__file__).resolve().parents[2]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--before-source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--format", choices=["xlsx", "pptx", "docx"], default="xlsx")
    parser.add_argument(
        "--plain-xlsx",
        action="store_true",
        help="Complete integer, boolean and text columns",
    )
    args = parser.parse_args()
    if args.plain_xlsx and args.format != "xlsx":
        parser.error("--plain-xlsx requires --format xlsx")
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    records = []
    unit = {"xlsx": "rows", "pptx": "slides", "docx": "paragraphs"}[args.format]
    sizes = {
        "xlsx": [10, 100, 1000, 10000],
        "pptx": [10, 100, 500],
        "docx": [10, 100, 1000, 5000],
    }[args.format]
    for size in sizes:
        path = output / f"{unit}-{size}.{args.format}"
        if args.format == "xlsx":
            from openpyxl import Workbook

            book = Workbook(write_only=True)
            sheet = book.create_sheet("Data")
            sheet.append([f"Field {i}" for i in range(10)])
            for n in range(size):
                sheet.append(
                    [
                        n,
                        n * 2 if args.plain_xlsx else n / 7,
                        f"项目 {n}",
                        "*a* & <b>",
                        "text" if args.plain_xlsx else None,
                        True,
                        "word" if args.plain_xlsx else "NA",
                        "data" if args.plain_xlsx else "001",
                        f"row {n}" if args.plain_xlsx else f"row\n{n}",
                        n % 3,
                    ]
                )
            book.save(path)
        elif args.format == "pptx":
            from pptx import Presentation

            deck = Presentation()
            for n in range(size):
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                title = slide.shapes.title
                assert title is not None
                title.text = f"Slide {n}"
                body = cast(Any, slide.placeholders[1])
                body.text = f"项目 {n}\nComplete body\nSecond paragraph"
            deck.save(str(path))
        else:
            template = REPO / "packages/markitai/tests/fixtures/sample.docx"
            with (
                zipfile.ZipFile(template) as source_archive,
                zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as target,
            ):
                for member in source_archive.namelist():
                    data = source_archive.read(member)
                    if member == "word/document.xml":
                        prefix, rest = data.decode().split("<w:body>", 1)
                        _, suffix = rest.split("</w:body>", 1)
                        paragraphs = "".join(
                            f"<w:p><w:r><w:rPr><w:b/></w:rPr>"
                            f"<w:t>Paragraph {n} 中文 &amp; complete body.</w:t>"
                            f"</w:r></w:p>"
                            for n in range(size)
                        )
                        data = (
                            prefix + "<w:body>" + paragraphs + "</w:body>" + suffix
                        ).encode()
                    target.writestr(member, data)
        for repeat in range(3):
            sources = [
                ("before", args.before_source),
                ("after", REPO / "packages/markitai/src"),
            ]
            for label, source in sources[:: 1 if repeat % 2 == 0 else -1]:
                env = {
                    **os.environ,
                    "PYTHONPATH": str(source.resolve()),
                    "LITELLM_LOCAL_MODEL_COST_MAP": "True",
                    "MARKITAI_NO_REMOTE_FETCH": "1",
                }
                command = [
                    sys.executable,
                    str(REPO / "scripts/benchmarks/audit_performance.py"),
                    "--worker",
                    "converter",
                    "--input",
                    str(path),
                    "--output",
                    str(output / "work"),
                ]
                start = time.perf_counter()
                result = subprocess.run(
                    command,
                    check=False,
                    capture_output=True,
                    text=True,
                    env=env,
                    timeout=120,
                )
                seconds = time.perf_counter() - start
                stem = f"{size}-{label}-{repeat}"
                (output / f"{stem}.stdout").write_text(result.stdout)
                (output / f"{stem}.stderr").write_text(result.stderr)
                assert result.returncode == 0, stem
                payload = next(
                    json.loads(line.removeprefix("AUDIT_RESULT="))
                    for line in result.stdout.splitlines()
                    if line.startswith("AUDIT_RESULT=")
                )
                records.append(
                    {
                        unit: size,
                        "size": size,
                        "engine": label,
                        "repeat": repeat,
                        "process_seconds": seconds,
                        **payload,
                    }
                )
                (output / "scaling.json").write_text(json.dumps(records, indent=2))
                print(stem, round(seconds, 3), flush=True)
        hashes = {
            i["sha256"] for r in records if r["size"] == size for i in r["iterations"]
        }
        assert len(hashes) == 1, f"Output mismatch at {size} {unit}"


if __name__ == "__main__":
    main()
